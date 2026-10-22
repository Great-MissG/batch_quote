from datetime import date
from io import BytesIO
from pathlib import Path
import unittest
import openpyxl
from quote_service import Shipment, zip5, cheapest_carriers, request_quote, read_upload, export_results, display_fields


def upload_bytes():
    wb = openpyxl.load_workbook(Path(__file__).parent / '询价模板.xlsx')
    wb.active.append([2, 48, 40, 36, 150, '08859-1234', '10001-9876'])
    wb.active.append([1, 48, 40, 36, 0, '91752', '10001'])
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


class QuoteTests(unittest.TestCase):
    def test_zip(self):
        self.assertEqual(zip5('08859-1234'), '08859')
        self.assertEqual(zip5(8859), '08859')
        for value in ['12345garbage', '123456789', '1234', '', None, 12345.7, True]:
            with self.assertRaises(ValueError): zip5(value)

    def test_units_and_payload(self):
        s = Shipment.parse(2, [2, 48, 40, 36, 150, '08859-1234', '10001'])
        calls = []
        def send(token, endpoint, payload):
            calls.append((endpoint, payload))
            if endpoint.endswith('suggested-class'): return {'suggestedClass': '100'}
            return {'options': [{'carrierCode': 'A', 'carrierName': 'Carrier A', 'totalAmount': 99.5}], 'quoteNo': 'TEST'}
        result = request_quote(s, '2026-10-22', 'test-only', send)
        for _, payload in calls:
            item = payload.get('items', [payload])[0]
            self.assertEqual(item['units'], 2)
            self.assertEqual(item['totalWeight'], 300)
        self.assertEqual(calls[1][1]['accessorialServices'], [])
        self.assertEqual(calls[1][1]['originZipCode'], '08859')
        item = calls[1][1]['items'][0]
        self.assertEqual(item['packagingType'], 'Pallet')
        for flag in ['isStackable', 'isHazardous', 'isUsed', 'isMachinery']:
            self.assertIs(item[flag], False)
        self.assertEqual(result['options'][0]['totalAmount'], 99.5)
        for qty in [0, -1, 1.5]:
            with self.assertRaises(ValueError): Shipment.parse(2, [qty, 48, 40, 36, 150, '08859', '10001'])

    def test_distinct_four_and_missing_fields(self):
        options = [{'carrierCode': c, 'totalAmount': p} for c, p in [('A', 200), ('a', 100), ('B', 80), ('C', 90), ('D', 110), ('E', 120)]]
        top = cheapest_carriers(options)
        self.assertEqual([o['totalAmount'] for o in top], [80, 90, 100, 110])
        self.assertEqual(len(top), 4)
        self.assertEqual(display_fields({'carrierCode': 'A', 'totalAmount': 0, 'platformInsuranceAmount': 0})[1:], ['未明确', 'A', '未提供'])
        self.assertEqual(display_fields({'carrierCode': 'A', 'totalAmount': 10, 'transitDays': 0})[3], '预计0天')

    def test_template_export_and_failure(self):
        data = upload_bytes()
        wb, shipments, errors = read_upload(data)
        self.assertEqual(len(shipments), 1)
        self.assertIn(3, errors)
        result = {2: {'options': [{'carrierCode': 'A', 'carrierName': '=UNTRUSTED()', 'totalAmount': 123.45, 'transitDays': 3}], 'quoteNo': 'TEST'}}
        output = export_results(data, result, errors)
        check = openpyxl.load_workbook(BytesIO(output))
        sh = check.active
        self.assertEqual(sh.max_column, 23)
        self.assertEqual(sh['H2'].value, 123.45)
        self.assertEqual(sh['J2'].data_type, 's')
        self.assertEqual(sh['I2'].value, '未明确')
        self.assertIsNone(sh['L2'].value)
        self.assertEqual(sh['H3'].value, '询价失败')
        self.assertEqual(sh['F2'].value, '08859-1234')
        self.assertEqual([c.value for c in sh[1]], [c.value for c in wb.active[1]])

    def test_upload_with_only_seven_input_columns(self):
        wb = openpyxl.load_workbook(BytesIO(upload_bytes()))
        wb.active.delete_cols(8, 16)
        stream = BytesIO()
        wb.save(stream)
        content = stream.getvalue()
        _, shipments, errors = read_upload(content)
        self.assertEqual(len(shipments), 1)
        self.assertIn(3, errors)
        output = export_results(content, {}, {})
        result = openpyxl.load_workbook(BytesIO(output)).active
        self.assertEqual(result.max_column, 23)
        from quote_service import HEADERS
        self.assertEqual([c.value for c in result[1]], HEADERS)
        self.assertEqual(result['F2'].value, '08859-1234')


if __name__ == '__main__':
    unittest.main()
