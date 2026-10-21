"""Fmile batch quotation. No order creation or automatic POST retries."""
from copy import copy
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
import json
import math
import re
import urllib.error
import urllib.request

import openpyxl

INPUT_HEADERS = ['托盘数量', '长（inch）', '宽（inch）', '高（inch）', '单一托盘重量（lbs）', '提货邮编', '收货邮编']
OUTPUT_HEADERS = [label for i in range(1, 5) for label in [f'报价{i}', '是否包含保险', '选择卡司' if i == 1 else f'选择卡司{i}', '预估时效']]
HEADERS = INPUT_HEADERS + OUTPUT_HEADERS
API_BASE = 'https://oms.fimile.com/oms-manager/api/ltl'


class QuoteError(Exception):
    pass


def zip5(value):
    if isinstance(value, bool) or value is None:
        raise ValueError('邮编不能为空')
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or value != int(value):
            raise ValueError('邮编必须是5位或ZIP+4')
        value = str(int(value)).zfill(5)
    text = str(value).strip()
    if not re.fullmatch(r'\d{5}(?:-\d{4})?', text, flags=re.ASCII):
        raise ValueError('邮编必须是5位或12345-6789格式')
    return text[:5]


def positive(value, label):
    if isinstance(value, bool) or value is None:
        raise ValueError(f'{label}必须大于0')
    try:
        result = Decimal(str(value).strip())
    except InvalidOperation:
        raise ValueError(f'{label}必须是数字') from None
    if not result.is_finite() or result <= 0:
        raise ValueError(f'{label}必须大于0')
    return result


@dataclass
class Shipment:
    row: int
    units: int
    length: float
    width: float
    height: float
    total_weight: float
    origin: str
    destination: str

    @classmethod
    def parse(cls, row, values):
        qty = positive(values[0], '托盘数量')
        if qty != qty.to_integral_value():
            raise ValueError('托盘数量必须为整数')
        dimensions = [float(positive(values[i], INPUT_HEADERS[i])) for i in range(1, 4)]
        weight = float((qty * positive(values[4], INPUT_HEADERS[4])).quantize(Decimal('0.01')))
        if not all(math.isfinite(v) for v in [*dimensions, weight]):
            raise ValueError('尺寸或重量数值过大')
        if weight <= 20:
            raise ValueError('整票总重量必须大于20 lbs')
        return cls(row, int(qty), *dimensions, weight, zip5(values[5]), zip5(values[6]))


def read_upload(content):
    if len(content) > 10 * 1024 * 1024:
        raise ValueError('文件不能超过10 MB')
    try:
        wb = openpyxl.load_workbook(BytesIO(content))
    except Exception:
        raise ValueError('无法读取文件，请上传有效的.xlsx工作簿') from None
    sh = wb.worksheets[0]
    for col, expected in enumerate(INPUT_HEADERS, 1):
        actual = sh.cell(1, col).value
        if str(actual or '').strip().lstrip('\ufeff') != expected:
            letter = openpyxl.utils.get_column_letter(col)
            raise ValueError(f'{letter}列应为“{expected}”，实际为“{actual or "空白"}”')
    for col, expected in enumerate(OUTPUT_HEADERS, 8):
        actual = sh.cell(1, col).value
        if actual is not None and str(actual).strip() not in ('', expected):
            letter = openpyxl.utils.get_column_letter(col)
            raise ValueError(f'{letter}列是结果区域，应为空白或“{expected}”，实际为“{actual}”')
    shipments, errors = [], {}
    for n, row in enumerate(sh.iter_rows(min_row=2, max_col=7, values_only=True), 2):
        if all(v is None or str(v).strip() == '' for v in row):
            continue
        try:
            shipments.append(Shipment.parse(n, row))
        except (ValueError, OverflowError) as exc:
            errors[n] = str(exc)
        if len(shipments) + len(errors) > 500:
            raise ValueError('一次最多处理500票，请拆分上传')
    return wb, shipments, errors


def post(token, endpoint, payload):
    req = urllib.request.Request(API_BASE + endpoint, data=json.dumps(payload).encode('utf-8'),
        headers={'Authorization': 'Bearer ' + token, 'Content-Type': 'application/json', 'Accept': 'application/json'}, method='POST')
    try:
        with urllib.request.urlopen(req, timeout=100) as response:
            return json.load(response)  # Python preserves 64-bit integer IDs.
    except urllib.error.HTTPError as exc:
        try:
            detail = json.loads(exc.read()).get('error', '接口请求失败')
        except (ValueError, AttributeError):
            detail = '接口请求失败'
        detail = str(detail).replace(token, '[已隐藏]')[:400]
        raise QuoteError(f'HTTP {exc.code}：{detail}') from None
    except (urllib.error.URLError, TimeoutError, OSError):
        raise QuoteError('网络连接失败或超时，未自动重试，请稍后手动重新询价') from None
    except (ValueError, TypeError):
        raise QuoteError('接口返回内容格式不正确') from None


def cheapest_carriers(options):
    carriers = {}
    for option in options:
        if not isinstance(option, dict):
            continue
        amount = option.get('totalAmount')
        if isinstance(amount, bool) or not isinstance(amount, (int, float)) or not math.isfinite(amount) or amount < 0:
            continue
        code = str(option.get('carrierCode') or option.get('carrierName') or '').strip().casefold()
        if code and (code not in carriers or amount < carriers[code]['totalAmount']):
            carriers[code] = option
    return sorted(carriers.values(), key=lambda o: o['totalAmount'])[:4]


def request_quote(shipment, pickup_date, token, send=post):
    dims = {'length': shipment.length, 'width': shipment.width, 'height': shipment.height}
    suggested = send(token, '/quotes/suggested-class', {**dims, 'totalWeight': shipment.total_weight, 'units': shipment.units})
    if not isinstance(suggested, dict):
        raise QuoteError('建议货运等级返回格式不正确')
    freight_class = suggested.get('suggestedClass')
    if not freight_class:
        raise QuoteError('接口未返回建议货运等级')
    payload = {'originZipCode': shipment.origin, 'destinationZipCode': shipment.destination,
        'originCountryCode': 'US', 'destinationCountryCode': 'US', 'pickupDate': pickup_date,
        'accessorialServices': [], 'items': [{**dims, 'freightClass': str(freight_class),
        'packagingType': 'PALLET', 'units': shipment.units, 'totalWeight': shipment.total_weight}]}
    response = send(token, '/quotes', payload)
    if not isinstance(response, dict) or not isinstance(response.get('options'), list):
        raise QuoteError('接口未返回有效的报价列表')
    choices = cheapest_carriers(response['options'])
    if not choices:
        raise QuoteError('无可用报价（NO_RATE或渠道不支持）')
    return {'options': choices, 'quoteNo': str(response.get('quoteNo', '')), 'freightClass': str(freight_class)}


def display_fields(option):
    days = option.get('transitDays')
    eta = option.get('estimatedDeliveryDate')
    timing = f'预计{days}天' if days is not None else '未提供'
    if eta:
        timing = f'{timing}；预计送达{eta}' if days is not None else f'预计送达{eta}'
    # A zero insurance amount is not evidence of coverage or lack of carrier liability.
    insurance = '未明确'
    return [option['totalAmount'], insurance, option.get('carrierName') or option.get('carrierCode'), timing]


def export_results(content, results, errors):
    wb, _, _ = read_upload(content)
    sh = wb.worksheets[0]
    template = openpyxl.load_workbook(Path(__file__).parent / '询价模板.xlsx')
    for col, label in enumerate(OUTPUT_HEADERS, 8):
        sh.cell(1, col, label)._style = copy(template.worksheets[0].cell(1, col)._style)
    sh.row_dimensions[1].height = max(sh.row_dimensions[1].height or 15, 32)
    for row in range(2, sh.max_row + 1):
        for col in range(8, 24):
            sh.cell(row, col).value = None
    for row, result in results.items():
        for i, option in enumerate(result['options']):
            for j, value in enumerate(display_fields(option)):
                cell = sh.cell(row, 8 + 4 * i + j, value)
                cell._style = copy(sh.cell(2, 8 + 4 * i + j)._style)
                if j == 0:
                    cell.number_format = '"$"#,##0.00'
                if isinstance(value, str):
                    cell.data_type = 's'  # Keep supplier text from becoming an Excel formula.
                cell.alignment = openpyxl.styles.Alignment(vertical='center', wrap_text=True)
        sh.row_dimensions[row].height = 42
    for row, message in errors.items():
        sh.cell(row, 8, '询价失败')
        sh.cell(row, 9, str(message)).data_type = 's'
        sh.cell(row, 9).alignment = openpyxl.styles.Alignment(wrap_text=True)
        sh.row_dimensions[row].height = 42
    for col in range(8, 24):
        letter = openpyxl.utils.get_column_letter(col)
        sh.column_dimensions[letter].width = 30 if (col - 8) % 4 == 3 else 20
    sh.freeze_panes = 'H2'
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
